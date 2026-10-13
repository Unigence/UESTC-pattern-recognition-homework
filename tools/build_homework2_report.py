"""Build the report layout and MathType insertion manifest from the Markdown source.

Requires python-docx and Pillow. Run in a document-build environment, then run
homework2_mathtype.ps1 to replace equation markers using installed MathType/Word.
The experimental notebook has its own independent, pinned environment.
"""
from pathlib import Path
from copy import deepcopy
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'report'
SOURCE = REPORT / 'homework2_report.md'
OUT = REPORT / '模式识别第二次作业-实验报告.docx'
BUILD = ROOT / 'tmp' / 'homework2_report'
BUILD.mkdir(parents=True, exist_ok=True)

def font(run, size=12, bold=False, east='SimSun'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)

def text_runs(p, text, size=12, bold=False, east='SimSun'):
    # Preserve URLs as readable source text; no source code enters the report.
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', r'\1（\2）', text)
    text = text.replace('`', '')
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            font(p.add_run(part), size, bold or i % 2 == 1, east)

def pformat(p, indent=True):
    pf = p.paragraph_format
    pf.first_line_indent = Pt(24) if indent else Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(20)
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def table(doc, lines, keep_with_note=False):
    rows = [[s.strip() for s in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', x) for x in r)]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # Use a wider text column and proportional numeric columns.
    count = len(rows[0])
    weights = [1.6] + [1] * (count - 1)
    if count == 2:
        weights = [1.0, 2.6]
    if count == 5 and rows[0][-1] == 'AUC':
        weights = [1.35, 1.05, 1.05, 1.05, 1.45]
    if count == 5 and rows[0][-1] == 'ΔAUC':
        weights = [0.5, 1.15, 1.15, 1.15, 1.0]
    widths = [14.65 * w / sum(weights) for w in weights]
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single' if edge in ['top', 'bottom'] else 'nil')
        if edge in ['top', 'bottom']:
            b.set(qn('w:sz'), '10')
            b.set(qn('w:color'), '000000')
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(rows):
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            pformat(p, False)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = i < len(rows) - 1 or keep_with_note
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_runs(p, value, 10.5, i == 0)
            if i == 0:
                cb = OxmlElement('w:tcBorders')
                bottom = OxmlElement('w:bottom')
                for attr, val in [('val', 'single'), ('sz', '6'), ('color', '000000')]:
                    bottom.set(qn('w:' + attr), val)
                cb.append(bottom)
                c._tc.get_or_add_tcPr().append(cb)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.keep_with_next = keep_with_note

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.175), Cm(3.175)
    sec.header_distance, sec.footer_distance = Cm(1.5), Cm(1.75)
    sec.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'SimSun')
    # Leave the default line height automatic so the inherited 36 pt cover
    # title is not clipped; body paragraphs receive their own 20 pt spacing.
    template = REPORT / '模式识别第一次作业-实验报告.docx'
    if template.exists():
        old = Document(template)
        doc.styles.element[:] = [deepcopy(node) for node in old.styles.element]
        for p in old.paragraphs[:22]:
            doc._body._body.insert(len(doc._body._body) - 1, deepcopy(p._p))
    else:
        for text in ['电子科技大学 自动化工程学院', '标准实验报告', '模式识别与机器学习', '第二次作业']:
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    foot._p.append(fld)
    equations = []
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            doc.add_page_break()
        elif line.startswith('<!--'):
            continue
        elif line.startswith('# '):
            p = doc.add_paragraph()
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            text_runs(p, line[2:], 15, True, 'SimHei')
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 1')
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            text_runs(p, line[3:], 14, True, 'SimHei')
        elif line.startswith('### '):
            p = doc.add_paragraph(style='Heading 2')
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            text_runs(p, line[4:], 12, True, 'SimHei')
        elif line.startswith('$$'):
            tex = line[2:]
            while not tex.endswith('$$'):
                tex += ' ' + lines[i].strip()
                i += 1
            tex = tex[:-2].strip()
            marker = f'EQ{len(equations)+1:03d}'
            equations.append({'id': marker, 'tex': tex})
            p = doc.add_paragraph(marker)
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            p.paragraph_format.line_spacing = Pt(24)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            next_text = next((x.strip() for x in lines[i:] if x.strip()), '')
            p.paragraph_format.keep_with_next = next_text.startswith('$$')
        elif line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            next_text = next((x.strip() for x in lines[i:] if x.strip()), '')
            table(doc, block, next_text.startswith('注：'))
        elif line.startswith('!['):
            alt, path = re.match(r'!\[(.*?)\]\((.*?)\)', line).groups()
            p = doc.add_paragraph()
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((REPORT / path).resolve()), width=Cm(14.65))
            p = doc.add_paragraph()
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_runs(p, alt, 10.5)
        elif re.match(r'^表\d+-\d+\s', line):
            p = doc.add_paragraph()
            pformat(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            text_runs(p, line, 10.5)
        elif line.startswith('注：'):
            p = doc.add_paragraph()
            pformat(p, False)
            text_runs(p, line, 10.5)
        else:
            p = doc.add_paragraph()
            pformat(p, not line.startswith('['))
            if line.startswith('['):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = Pt(15)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
            if line.endswith('：'):
                p.paragraph_format.keep_with_next = True
            text_runs(p, line, 10.5 if line.startswith('[') else 12)
    doc.core_properties.title = '基于 BP 神经网络与支持向量机的男女生分类'
    doc.core_properties.subject = '模式识别第二次作业：五折交叉验证与结果分析'
    doc.core_properties.author = '周子智'
    doc.save(OUT)
    manifest = BUILD / 'equations.json'
    manifest.write_text(json.dumps(equations, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Saved {OUT}; {len(equations)} equations in {manifest}')

if __name__ == '__main__':
    main()
